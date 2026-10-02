"""Phase 7 — Part 3d: documents + learner."""

from __future__ import annotations

import json
import zipfile
from pathlib import Path

import pytest


# ---------------------------------------------------------------------------
# Documents
# ---------------------------------------------------------------------------
def test_documents_available_reports_all_three():
    from core.documents import available
    out = available()
    assert out["docx"] is True
    assert out["xlsx"] is True
    assert out["pptx"] is True


def test_create_report_writes_valid_docx(tmp_path):
    from core.documents import create_report
    res = create_report(
        title="Phase 7 report",
        sections=[
            {"heading": "Summary",
              "body": "This is a paragraph.\n\nAnd another."},
            {"heading": "Findings",
              "body": ["First bullet", "Second bullet"]},
        ],
        directory=tmp_path,
    )
    assert res.kind == "docx"
    p = Path(res.path)
    assert p.exists()
    # docx is a zip containing word/document.xml.
    with zipfile.ZipFile(p) as zf:
        names = zf.namelist()
        assert "word/document.xml" in names
        body = zf.read("word/document.xml").decode("utf-8", "replace")
    assert "Phase 7 report" in body
    assert "First bullet" in body


def test_create_spreadsheet_writes_valid_xlsx(tmp_path):
    from core.documents import create_spreadsheet
    res = create_spreadsheet(
        title="Findings",
        sheets=[
            {"name": "Tools",
              "headers": ["Name", "Score"],
              "rows": [["nmap", 9], ["masscan", 7]]},
        ],
        directory=tmp_path,
    )
    p = Path(res.path)
    assert p.exists()
    from openpyxl import load_workbook
    wb = load_workbook(p)
    ws = wb["Tools"]
    assert ws.cell(1, 1).value == "Name"
    assert ws.cell(2, 1).value == "nmap"
    assert ws.cell(2, 2).value == 9


def test_create_presentation_writes_valid_pptx(tmp_path):
    from core.documents import create_presentation
    res = create_presentation(
        title="Quarterly review",
        slides=[
            {"title": "Highlights",
              "bullets": ["Shipped Phase 6", "Started Phase 7"]},
        ],
        directory=tmp_path,
    )
    p = Path(res.path)
    assert p.exists()
    with zipfile.ZipFile(p) as zf:
        names = zf.namelist()
        assert any(n.startswith("ppt/slides/") for n in names)


# ---------------------------------------------------------------------------
# Learner
# ---------------------------------------------------------------------------
@pytest.fixture
def fake_skills_dir(tmp_path, monkeypatch):
    """Redirect SKILLS_DIR so the learner doesn't write into the repo."""
    from core import learner, paths
    fake_root = tmp_path / "skills"
    fake_root.mkdir()
    monkeypatch.setattr(paths, "SKILLS_DIR", fake_root)
    monkeypatch.setattr(learner, "SKILLS_DIR", fake_root)
    return fake_root


def test_learner_rejects_missing_goal(fake_skills_dir):
    from core.learner import learn_from_demonstration
    with pytest.raises(ValueError):
        learn_from_demonstration({"steps": ["do a thing"]})


def test_learner_rejects_missing_steps(fake_skills_dir):
    from core.learner import learn_from_demonstration
    with pytest.raises(ValueError):
        learn_from_demonstration({"goal": "do something"})


def test_learner_persists_skill_md_and_demo_json(fake_skills_dir):
    from core.learner import learn_from_demonstration, list_learned
    learned = learn_from_demonstration({
        "goal": "fetch and summarise CVE data",
        "steps": ["GET https://nvd.nist.gov/...", "summarise"],
        "inputs": {"query": "log4shell"},
        "tags": ["security"],
    })
    folder = Path(learned.path)
    skill_md = folder / "SKILL.md"
    demo_json = folder / "demo.json"
    assert skill_md.exists()
    assert demo_json.exists()
    text = skill_md.read_text()
    assert "fetch and summarise CVE data" in text
    assert "<query>" in text
    demo = json.loads(demo_json.read_text())
    assert demo["inputs"]["query"] == "log4shell"

    # list_learned reflects the new entry.
    out = list_learned()
    assert any(s["slug"] == learned.slug for s in out)


def test_learner_marks_generalization_skipped_without_llm(
        fake_skills_dir, monkeypatch):
    """When no LLM is available we keep the demo verbatim."""
    from core import learner
    monkeypatch.setattr(learner, "_llm_generalize", lambda demo: None)
    learned = learner.learn_from_demonstration({
        "goal": "x", "steps": ["a", "b"]})
    assert learned.generalization_skipped is True


def test_learner_uses_llm_when_available(fake_skills_dir, monkeypatch):
    from core import learner

    def fake_llm(demo):
        return {"title": "Generalised", "summary": "summary",
                 "steps": ["step1", "step2"],
                 "generalization_skipped": False}

    monkeypatch.setattr(learner, "_llm_generalize", fake_llm)
    learned = learner.learn_from_demonstration({
        "goal": "do thing", "steps": ["raw"]})
    assert learned.generalization_skipped is False
    assert learned.title == "Generalised"
    assert "step1" in learned.steps


def test_load_learned_returns_skill_md(fake_skills_dir):
    from core.learner import learn_from_demonstration, load_learned
    learned = learn_from_demonstration({
        "goal": "x", "steps": ["a"]})
    out = load_learned(learned.slug)
    assert out is not None
    assert "x" in out["skill_md"]


# ---------------------------------------------------------------------------
# API endpoints
# ---------------------------------------------------------------------------
def test_api_documents_endpoints(monkeypatch, tmp_path):
    from fastapi.testclient import TestClient
    from api.server import app
    monkeypatch.setenv("AGENT_OS_API_TOKEN", "secret")
    # Redirect WORKSPACE_DIR so generated files land in tmp_path.
    from core import paths, documents
    monkeypatch.setattr(paths, "WORKSPACE_DIR", tmp_path)
    monkeypatch.setattr(documents, "WORKSPACE_DIR", tmp_path)
    client = TestClient(app)
    h = {"X-Agent-Token": "secret"}

    r = client.get("/admin/documents/available", headers=h)
    assert r.status_code == 200
    assert r.json()["docx"] is True

    r = client.post("/admin/documents/report", headers=h, json={
        "title": "Test", "sections": [{"heading": "h", "body": "b"}]})
    assert r.status_code == 200
    assert r.json()["kind"] == "docx"

    r = client.post("/admin/documents/spreadsheet", headers=h, json={
        "title": "S", "sheets": [{"name": "S1",
                                       "headers": ["a"], "rows": [[1]]}]})
    assert r.status_code == 200

    r = client.post("/admin/documents/presentation", headers=h, json={
        "title": "P", "slides": [{"title": "Intro",
                                       "bullets": ["x"]}]})
    assert r.status_code == 200


def test_api_skills_learn_endpoint(monkeypatch, tmp_path):
    from fastapi.testclient import TestClient
    from api.server import app
    from core import learner, paths
    monkeypatch.setenv("AGENT_OS_API_TOKEN", "secret")
    fake_root = tmp_path / "skills"
    fake_root.mkdir()
    monkeypatch.setattr(paths, "SKILLS_DIR", fake_root)
    monkeypatch.setattr(learner, "SKILLS_DIR", fake_root)
    client = TestClient(app)
    h = {"X-Agent-Token": "secret"}

    r = client.post("/admin/skills/learn_from_demonstration",
                     headers=h, json={"goal": "g", "steps": ["s"]})
    assert r.status_code == 200, r.text
    assert "slug" in r.json()

    r = client.get("/admin/skills/learned", headers=h)
    assert r.status_code == 200
    assert len(r.json()["skills"]) >= 1

    # Bad payload → 400.
    r = client.post("/admin/skills/learn_from_demonstration",
                     headers=h, json={"steps": []})
    assert r.status_code == 400
